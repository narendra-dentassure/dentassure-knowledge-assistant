"""Load PDF, DOCX, and TXT documents from a local folder."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from src.logging_setup import configure_logging

logger = configure_logging()

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}
# Character budget for DOCX "pages" so citations stay page-like, not one giant blob.
DOCX_SECTION_CHARS = 900


@dataclass
class LoadedDocument:
    """A source document split into page-like sections."""

    source: str
    file_type: str
    pages: list[tuple[int, str]] = field(default_factory=list)


def _read_txt(path: Path) -> list[tuple[int, str]]:
    text = path.read_text(encoding="utf-8", errors="replace").strip()
    if not text:
        return []
    # Keep TXT files as numbered sections so citations stay useful.
    sections = [part.strip() for part in text.split("\n\n") if part.strip()]
    if not sections:
        return [(1, text)]
    return [(index + 1, section) for index, section in enumerate(sections)]


def _read_pdf(path: Path) -> list[tuple[int, str]]:
    reader = PdfReader(str(path))
    pages: list[tuple[int, str]] = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append((index, text))
    return pages


def _read_docx(path: Path) -> list[tuple[int, str]]:
    document = Document(str(path))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    if not paragraphs:
        return []

    pages: list[tuple[int, str]] = []
    buffer: list[str] = []
    page_number = 1
    for paragraph in paragraphs:
        buffer.append(paragraph)
        if len(" ".join(buffer)) >= DOCX_SECTION_CHARS:
            pages.append((page_number, "\n".join(buffer)))
            buffer = []
            page_number += 1
    if buffer:
        pages.append((page_number, "\n".join(buffer)))
    return pages


def load_documents(data_dir: Path) -> list[LoadedDocument]:
    """Load all supported documents from data_dir."""
    if not data_dir.exists():
        raise FileNotFoundError(f"Sample data folder not found: {data_dir}")

    loaded: list[LoadedDocument] = []
    files = sorted(
        path
        for path in data_dir.rglob("*")
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS
    )
    if not files:
        raise FileNotFoundError(f"No PDF, DOCX, or TXT files found in {data_dir}")

    for path in files:
        suffix = path.suffix.lower()
        try:
            if suffix == ".pdf":
                pages = _read_pdf(path)
            elif suffix == ".docx":
                pages = _read_docx(path)
            else:
                pages = _read_txt(path)
        except Exception:
            logger.exception("Failed to read document: %s", path.name)
            continue

        if not pages:
            logger.warning("Skipping empty document: %s", path.name)
            continue

        loaded.append(
            LoadedDocument(source=path.name, file_type=suffix.lstrip("."), pages=pages)
        )
        logger.info("Loaded %s (%s pages/sections)", path.name, len(pages))

    if not loaded:
        raise RuntimeError("All documents failed to load or were empty.")
    return loaded
