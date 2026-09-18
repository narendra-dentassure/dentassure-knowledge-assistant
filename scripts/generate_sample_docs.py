"""Generate sample PDF, DOCX, and TXT documents."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fpdf import FPDF

ROOT = Path(__file__).resolve().parent.parent
OUTPUT_DIR = ROOT / "data" / "sample_docs"


class HandbookPDF(FPDF):
    """Simple multi-page PDF writer for sample knowledge documents."""

    def header(self) -> None:
        self.set_font("Helvetica", "B", 10)
        self.set_text_color(12, 74, 110)
        self.set_x(self.l_margin)
        self.cell(0, 8, self.title, align="L", new_x="LMARGIN", new_y="NEXT")
        self.set_draw_color(13, 148, 136)
        self.line(10, 16, 200, 16)
        self.ln(8)

    def footer(self) -> None:
        self.set_y(-15)
        self.set_font("Helvetica", "I", 8)
        self.set_text_color(100, 100, 100)
        self.cell(0, 10, f"Page {self.page_no()}", align="C")

    def section(self, heading: str, body: str) -> None:
        self.set_x(self.l_margin)
        self.set_font("Helvetica", "B", 13)
        self.set_text_color(15, 23, 42)
        self.multi_cell(0, 8, heading, new_x="LMARGIN", new_y="NEXT")
        self.ln(1)
        self.set_font("Helvetica", "", 11)
        self.set_text_color(30, 41, 59)
        self.multi_cell(0, 6, body, new_x="LMARGIN", new_y="NEXT")
        self.ln(4)


def _write_pdf(filename: str, title: str, sections: list[tuple[str, str]]) -> None:
    pdf = HandbookPDF()
    pdf.set_title(title)
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 18)
    pdf.multi_cell(0, 10, title, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 11)
    pdf.multi_cell(
        0,
        6,
        "Internal reference for DentAssure members, visitors, and network clinic staff.",
        new_x="LMARGIN",
        new_y="NEXT",
    )
    for heading, body in sections:
        pdf.add_page()
        pdf.section(heading, body)
    pdf.output(OUTPUT_DIR / filename)


def _write_docx(filename: str, title: str, sections: list[tuple[str, str]]) -> None:
    document = Document()
    heading = document.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    document.add_paragraph(
        "Front-desk SOP for DentAssure cashless at a network clinic."
    )
    for heading_text, body in sections:
        document.add_heading(heading_text, level=1)
        for paragraph in body.split("\n\n"):
            document.add_paragraph(paragraph.strip())
    document.save(OUTPUT_DIR / filename)


def _write_txt(filename: str, title: str, sections: list[tuple[str, str]]) -> None:
    lines = [title, "=" * len(title), ""]
    for heading, body in sections:
        lines.extend([heading, "-" * len(heading), body.strip(), ""])
    (OUTPUT_DIR / filename).write_text("\n".join(lines), encoding="utf-8")


POLICY_SECTIONS = [
    (
        "1. Plan Overview",
        "DentAssure Health Plans provide outpatient dental coverage for individuals, families, "
        "and employer groups across India. The current flagship products are DentAssure Care, "
        "DentAssure Plus, and DentAssure Corporate. Coverage applies only to treatments "
        "performed at DentAssure network clinics. "
        "Policy year is 12 months from the admission date printed on the member card.",
    ),
    (
        "2. Waiting Periods and Eligibility",
        "Preventive care such as oral examination, scaling, and fluoride is covered from day 1. "
        "Restorative procedures including fillings and simple extractions have a 30-day waiting "
        "period. Root canal treatment, crowns, and bridges have a 90-day waiting period. "
        "Implants, orthodontics, and cosmetic veneers are excluded unless the Corporate rider "
        "is active. Dependents can be added within 30 days of the primary member enrolment. "
        "Pre-existing dental conditions must be declared at enrolment; undeclared conditions "
        "may be declined at claim stage.",
    ),
    (
        "3. Policy Renewal Guidelines",
        "Renewal is due 30 days before policy expiry. Members receive an SMS and email reminder "
        "on T-30, T-15, and T-7. Grace period is 15 days after expiry. During grace, emergency "
        "pain relief and extraction remain covered, but planned crowns and elective RCT are paused. "
        "If premium is not received by the end of grace, the policy lapses and a fresh waiting "
        "period applies. No-claim bonus is 10 percent extra annual limit for DentAssure Plus "
        "and 5 percent for DentAssure Care, capped at 50 percent of the original annual limit. "
        "Corporate policies renew on the employer anniversary date, not the individual start date.",
    ),
    (
        "4. Annual Limits and Co-pay",
        "DentAssure Care annual limit is Rs 15,000 per member with 20 percent co-pay on restorative "
        "work. DentAssure Plus annual limit is Rs 40,000 per member with 10 percent co-pay. "
        "DentAssure Corporate annual limit is Rs 75,000 per employee, with zero co-pay for "
        "preventive visits and 10 percent co-pay for specialist procedures. Pediatric members "
        "under 14 years have no co-pay on fluoride, pit and fissure sealants, or space maintainers. "
        "Once the annual limit is exhausted, the clinic may continue treatment on a private-pay "
        "basis after written consent at the clinic.",
    ),
    (
        "5. Claims Submission Rules",
        "Network clinics must submit claims to DentAssure within 7 calendar days of treatment "
        "completion. Required attachments are FDI tooth chart, treatment notes, pre and post "
        "photographs for crowns, and the signed consent form. Cashless claims under Rs 8,000 "
        "are auto-adjudicated within 4 working hours. Claims above Rs 8,000 require desk review "
        "and are decided within 3 working days. Reimbursement claims from non-network clinics "
        "are accepted only for emergency trauma and paid at 70 percent of the network tariff. "
        "Duplicate claims for the same tooth and same CDT-equivalent procedure within 90 days "
        "are automatically rejected.",
    ),
    (
        "6. Pre-authorization",
        "Pre-authorization is mandatory for RCT with crown, implants, orthodontics, and any "
        "treatment estimate above Rs 12,000. The treating dentist submits the "
        "pre-auth pack: diagnosis, proposed visits, chair time, and lab prescription. "
        "DentAssure desk responds in 1 working day for Plus and Corporate, and 2 working days "
        "for Care. Treatment started without pre-auth, where pre-auth was required, is paid "
        "at 50 percent of the approved tariff. Emergency pain cases can start immediately; "
        "pre-auth must still be filed the next working day.",
    ),
    (
        "7. Exclusions",
        "The following are excluded from all DentAssure plans: purely cosmetic whitening, "
        "diamond jewelry on teeth, experimental laser bleaching, missed-appointment fees, "
        "and treatments started before the policy start date. Sports-guard replacement is "
        "limited to one every 24 months. Night guards are covered only when a TMJ diagnosis "
        "is recorded in the FDI chart. Tobacco-stain polishing is not classified as scaling "
        "and is therefore not covered.",
    ),
    (
        "8. Member Support",
        "Members can check remaining annual limit in the DentAssure app or by asking the "
        "clinic front desk. DentAssure Customer Care: +91 888 668 6850, "
        "support@dentassureplans.com, www.dentassureplans.co.in. "
        "Complaints about claim delay must be raised within 15 days of the claim decision. "
        "If a clinic is removed from the network, in-progress multi-visit treatments already "
        "pre-authorized remain payable until completion or 60 days, whichever is earlier.",
    ),
]

SOP_SECTIONS = [
    (
        "1. Who this SOP is for",
        "This SOP is for reception at a DentAssure network clinic. It explains how a visitor "
        "or member uses a DentAssure health plan for cashless dental care. It is not a "
        "separate product. DentAssure is the plan. The clinic is where treatment happens.",
    ),
    (
        "2. Before the patient sits in the chair",
        "Ask: member or new visitor? If a member, check the plan name, start date, and remaining "
        "annual limit. If a visitor, show Our Plans on dentassureplans.co.in or ask this assistant "
        "which plan fits the treatments they want. Do not quote a cashless amount until the plan "
        "is confirmed.",
    ),
    (
        "3. Cashless visit steps",
        "1. Confirm identity and plan. 2. Confirm the clinic is in-network. 3. List the planned "
        "treatments (consultation, scaling, RCT, implant). 4. If the estimate is above Rs 12,000, "
        "or includes RCT with crown or implants, raise pre-authorization. 5. After treatment, "
        "submit the claim within 7 calendar days with FDI chart and consent.",
    ),
    (
        "4. What reception should tell the patient",
        "Preventive care such as examination, scaling, and fluoride is covered from day 1. "
        "Fillings wait 30 days. RCT, crowns, and bridges wait 90 days unless the specific plan "
        "table says otherwise. Customer Care is +91 888 668 6850. Website: dentassureplans.co.in.",
    ),
    (
        "5. When to send the patient to compare-services",
        "If the patient does not have a plan yet and names treatments, compare those treatments "
        "across published plans (same flow as compare-services on the website). Recommend the "
        "lowest total out-of-pocket option and ask them to enroll before the visit if possible.",
    ),
]


MEMBER_FAQ_SECTIONS = [
    (
        "Who can use this assistant",
        "Anyone: a patient with a DentAssure plan, a person who is only browsing, or clinic "
        "front desk helping a walk-in. The assistant answers from DentAssure policy documents "
        "and public plan/clinic information. It does not replace a dentist.",
    ),
    (
        "How to take a plan",
        "Go to https://dentassureplans.co.in/our-plans or enroll at a network clinic after KYC. "
        "Pay the plan amount. Coverage is at in-network clinics. Support: +91 888 668 6850, "
        "support@dentassureplans.com.",
    ),
    (
        "First visit",
        "Carry photo ID and your member card or app. Ask reception to confirm cashless before "
        "treatment starts. Find clinics at https://dentassureplans.co.in/our-network-clinics "
        "or https://dentassureplans.co.in/all-clinics.",
    ),
    (
        "If something is not in the documents",
        "The assistant will say it does not know. Call DentAssure Customer Care "
        "+91 888 668 6850 or visit a network clinic. Do not rely on invented prices or medical advice.",
    ),
]


MEMBER_HANDBOOK_SECTIONS = [
    (
        "1. What DentAssure is",
        "DentAssure Health Plans is a dental membership and outpatient dental cover. "
        "You buy a plan, then take treatment at a network clinic. Public pages: "
        "dentassureplans.co.in, Our Plans, About Plans, network clinics, all clinics, and blog.",
    ),
    (
        "2. How to use your plan at a clinic",
        "Show the plan at reception. Reception checks waiting period, remaining limit, and "
        "whether pre-auth is needed. Cashless is only at in-network clinics. Emergency pain "
        "can start immediately; pre-auth is filed the next working day if required.",
    ),
    (
        "3. How plan compare works",
        "If you name treatments (for example OP consultation and scaling), DentAssure ranks "
        "published plans by treatment cost plus plan amount. That is the same idea as "
        "compare-services on the website. Informal English is fine: rct scaling konsa plan lena.",
    ),
]


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    _write_pdf(
        "dentassure_policy_guide.pdf",
        "DentAssure Policy Guide - Renewal, Claims, and Coverage",
        POLICY_SECTIONS,
    )
    _write_pdf(
        "dentassure_member_handbook.pdf",
        "How to use your DentAssure plan",
        MEMBER_HANDBOOK_SECTIONS,
    )
    _write_docx(
        "dentassure_clinic_front_desk_sop.docx",
        "DentAssure network clinic - front desk SOP",
        SOP_SECTIONS,
    )
    _write_txt(
        "dentassure_member_faq.txt",
        "DentAssure member and visitor FAQ",
        MEMBER_FAQ_SECTIONS,
    )
    print(f"Wrote generated sample documents to {OUTPUT_DIR}")
    print("Left in place (hand-authored): dentassure_plan_selection_guide.txt, dentassure_network_clinics.txt")


if __name__ == "__main__":
    main()
